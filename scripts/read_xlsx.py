# -*- coding: utf-8 -*-
"""Чтение листа «Лоты 12» из СНП ВКО.xlsx"""
import openpyxl, json

wb = openpyxl.load_workbook('/home/z/my-project/upload/СНП ВКО.xlsx', data_only=True)
print("Листы:", wb.sheetnames)
ws = wb['Лоты 12'] if 'Лоты 12' in wb.sheetnames else wb[wb.sheetnames[0]]
print("Размер:", ws.max_row, "x", ws.max_column)

for r in range(1, min(ws.max_row, 40) + 1):
    vals = []
    for c in range(1, min(ws.max_column, 12) + 1):
        v = ws.cell(row=r, column=c).value
        if v is not None:
            vals.append(f"[{c}] {v}")
    if vals:
        print(f"R{r}: " + " | ".join(vals))
