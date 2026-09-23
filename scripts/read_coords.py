#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Полные координаты из ссылки Google Earth"""
import openpyxl, re, json

wb = openpyxl.load_workbook('/home/z/my-project/upload/СНП ВКО.xlsx', data_only=True)
ws = wb['Лоты 12']

rows = []
header = [c.value for c in ws[1]]
print(header)
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[4] is None:
        continue
    url = str(row[5])
    # формат: @lat,lon,zoom...
    m = re.search(r'@(-?\d+\.\d+),\s*(-?\d+\.\d+)', url)
    lat, lon = float(m.group(1)), float(m.group(2))
    rec = {
        'kato': row[1], 'district': row[2], 'okrug': row[3],
        'name': str(row[4]).replace('с.', '').strip(),
        'lat': lat, 'lon': lon, 'households': int(row[6])
    }
    rows.append(rec)
    print(f"{rec['name']}: lat={lat}, lon={lon}, ДХ={rec['households']}")

with open('/home/z/my-project/work/snp_data.json', 'w', encoding='utf-8') as f:
    json.dump(rows, f, ensure_ascii=False, indent=2)
print("\nСохранено: work/snp_data.json")
