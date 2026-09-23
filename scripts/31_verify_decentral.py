# -*- coding: utf-8 -*-
"""
Шаг 31. Семантическая проверка листа «Децентрализация ОРШ» (после QA-пайплайна).

Сверяет значения листа (включая кэш recalc для формул) с источниками:
  - work/boq_decentral_data.json  — итоговый расчёт схемы D (шаг 29);
  - work/decentral_explore.json   — sweep + стоимость (шаг 28);
  - work/boq_data.json / boq_cascade_data.json / boq_cascade16_data.json — схемы A/B/C.

Проверяет: блоки Б/В/Г/Д1/Д2, данные графика, порядок листов, пункты «Методики»
31-36, неизменность остальных листов (спот-чеки «Сводной», «Параметров сетей»,
«Сравнения схем», «Проверки волокно-км»).
"""
import json, importlib.util
from openpyxl import load_workbook

BASE = '/home/z/my-project'
OUT = f'{BASE}/download/Сводная_таблица_материалов_FTTH_ВКО_каскад.xlsx'

DEXP = json.load(open(f'{BASE}/work/decentral_explore.json', encoding='utf-8'))
DBOOK = json.load(open(f'{BASE}/work/boq_decentral_data.json', encoding='utf-8'))
ABOOK = json.load(open(f'{BASE}/work/boq_data.json', encoding='utf-8'))
BBOOK = json.load(open(f'{BASE}/work/boq_cascade_data.json', encoding='utf-8'))
CBOOK = json.load(open(f'{BASE}/work/boq_cascade16_data.json', encoding='utf-8'))

spec = importlib.util.spec_from_file_location('de28', f'{BASE}/scripts/28_decentral_explore.py')
de28 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(de28)

wb = load_workbook(OUT, data_only=True)          # кэш recalc (формулы уже вычислены)
wbf = load_workbook(OUT)                          # формулы
ws = wb['Децентрализация ОРШ']
wsf = wbf['Децентрализация ОРШ']

fails = []
checked = 0


def chk(desc, actual, expected, tol=0.051):
    global checked
    checked += 1
    ok = (abs(actual - expected) <= tol) if isinstance(expected, (int, float)) \
        else (actual == expected)
    if not ok:
        fails.append(f'{desc}: лист={actual!r} vs ожидалось={expected!r}')


def chkf(desc, cell_ref, expected, tol=0.051):
    """Сверка значения формулы по кэшу recalc + наличие формулы."""
    global checked
    checked += 1
    val = ws[cell_ref].value
    frm = wsf[cell_ref].value
    ok = (abs(val - expected) <= tol) if isinstance(expected, (int, float)) \
        else (val == expected)
    if not ok or not (isinstance(frm, str) and frm.startswith('=')):
        fails.append(f'{desc} [{cell_ref}]: кэш={val!r}, формула={frm!r}, ожидалось={expected!r}')


# ---------- порядок листов ----------
ORDER = ['Сводная', 'Параметры сетей', 'Сравнение схем', 'Оптимум по длине',
         'Проверка волокно-км', 'Децентрализация ОРШ', 'Методика']
chk('порядок листов', wb.sheetnames, ORDER)

# ---------- блок Б: зоны (строки 16-49, итог 50) ----------
zi = 0
for v in DBOOK['villages']:
    for j, z in enumerate(v['zones']):
        r = 16 + zi
        chk(f'Б: {v["name"]} зона {j} ДХ', ws.cell(row=r, column=6).value, z['houses'], 0)
        chk(f'Б: {v["name"]} зона {j} удалённость', ws.cell(row=r, column=7).value,
            round(z['root_dist_m']), 0)
        chk(f'Б: {v["name"]} зона {j} сплиттеры', ws.cell(row=r, column=8).value, z['splitters'], 0)
        chk(f'Б: {v["name"]} зона {j} фидер', ws.cell(row=r, column=9).value, z['feeder_fibers'], 0)
        chk(f'Б: {v["name"]} зона {j} порты', ws.cell(row=r, column=10).value, z['orsh_ports'], 0)
        if j == 0:
            chk(f'Б: {v["name"]} ЦУ первым (тип)', ws.cell(row=r, column=5).value, 'ЦУ (корневая)')
            chk(f'Б: {v["name"]} ЦУ удалённость 0', ws.cell(row=r, column=7).value, 0, 0)
        zi += 1
chk('Б: всего строк зон', zi, 34, 0)
chkf('Б: ИТОГО ДХ', 'F50', 2334, 0)
chkf('Б: ИТОГО сплиттеры', 'H50', DBOOK['totals']['splitters64'], 0)
chkf('Б: ИТОГО порты', 'J50', DBOOK['totals']['orsh_ports'], 0)
chkf('Б: ИТОГО фидер', 'I50', sum(z['feeder_fibers'] for v in DBOOK['villages']
                                  for z in v['zones']), 0)

# ---------- блок В: sweep (строки 55-66) ----------
SW = DEXP['sweep']
for i, s in enumerate(SW):
    r = 55 + i
    chk(f'В: S_MIN={s["s_min"]} порог', ws.cell(row=r, column=3).value, s['s_min'], 0)
    chk(f'В: S_MIN={s["s_min"]} ОРШ', ws.cell(row=r, column=4).value, s['orsh'], 0)
    chk(f'В: S_MIN={s["s_min"]} волокно-км', ws.cell(row=r, column=5).value, s['fiber_km'], 0.51)
    chk(f'В: S_MIN={s["s_min"]} кабель', ws.cell(row=r, column=7).value, s['cable_km'], 0.51)
    chk(f'В: S_MIN={s["s_min"]} сплиттеры', ws.cell(row=r, column=10).value, s['splitters'], 0)
    chk(f'В: S_MIN={s["s_min"]} сварки', ws.cell(row=r, column=11).value, s['splices'], 0)
    chk(f'В: S_MIN={s["s_min"]} порты', ws.cell(row=r, column=12).value, s['orsh_ports'], 0)
    chk(f'В: S_MIN={s["s_min"]} ср.маршрут', ws.cell(row=r, column=13).value, s['avg_route'], 0.51)
    chk(f'В: S_MIN={s["s_min"]} макс.ёмкость', ws.cell(row=r, column=14).value, s['top_fibers'], 0)
    chkf(f'В: S_MIN={s["s_min"]} Δ к A', f'F{r}', s['fiber_km'] / 4375.1 - 1, 0.001)
    chkf(f'В: S_MIN={s["s_min"]} ВСЕГО', f'I{r}', s['cable_km'] + 95.8, 0.11)
# рекомендованная строка = данным итоговой книги
r15 = 55 + [s['s_min'] for s in SW].index(15)
chk('В: S_MIN=15 волокно-км = книга D', ws.cell(row=r15, column=5).value,
    DBOOK['totals']['fiber_km'], 0.51)
chk('В: S_MIN=15 ОРШ = книга D', ws.cell(row=r15, column=4).value, DBOOK['totals']['orsh'], 0)

# ---------- данные графика (строки 70-81) ----------
sorted_sw = sorted(SW, key=lambda x: x['orsh'])
for i, s in enumerate(sorted_sw):
    r = 70 + i
    chk(f'График: ОРШ', ws.cell(row=r, column=2).value, s['orsh'], 0)
    chk(f'График: волокно-км', ws.cell(row=r, column=3).value, s['fiber_km'], 0.51)
    chk(f'График: стоимость', ws.cell(row=r, column=4).value,
        DEXP['d_costs'][str(s['s_min'])]['total_with_boxes'], 0.06)

# ---------- блок Г: сравнение 4 схем (строки 85-108) ----------
mtA, mtB, mtC, mtD = (ABOOK['materials_total'], BBOOK['materials_total'],
                      CBOOK['materials_total'], DBOOK['materials_total'])
tB, tC, tD = BBOOK['totals'], CBOOK['totals'], DBOOK['totals']
cab_tot = lambda mt: round(sum(mt[f'cable_{s}'] for s in de28.STD_FIBERS), 1)
susB = mtB['suspend_kits'] + mtB.get('drop_ext_suspends', 0)
susC = mtC['suspend_kits'] + mtC.get('drop_ext_suspends', 0)
avg_drop = lambda book: round(sum(v.get('drop_km_total', v.get('drop_km'))
                                  for v in book['villages']) / 2334 * 1000, 1)
max_drop = lambda book: max(v['max_drop_m'] for v in book['villages'])

G_ROWS = [
    (85, 'ДХ', [2334, 2334, 2334, 2334], 0, None),
    (86, 'волокно-км', [4375.1, tB['fiber_km'], tC['fiber_km'], tD['fiber_km']], 0.51, 'C'),
    (88, 'магистраль', [cab_tot(mtA), cab_tot(mtB), cab_tot(mtC), cab_tot(mtD)], 0.51, 'C'),
    (89, '8F', [mtA['cable_8'], mtB['cable_8'], mtC['cable_8'], mtD['cable_8']], 0.51, None),
    (90, '96F', [mtA['cable_96'], mtB['cable_96'], mtC['cable_96'], mtD['cable_96']], 0.51, None),
    (91, 'дропы', [mtA['drop_cable_km'], mtB['drop_cable_km'], mtC['drop_cable_km'],
                   mtD['drop_cable_km']], 0.51, 'A'),
    (93, 'ср. дроп', [avg_drop(ABOOK), avg_drop(BBOOK), avg_drop(CBOOK), avg_drop(DBOOK)], 0.51, 'A'),
    (94, 'макс. дроп', [max_drop(ABOOK), max_drop(BBOOK), max_drop(CBOOK), max_drop(DBOOK)], 0.51, 'A'),
    (95, 'ср. маршрут', [1361, None, None,
                         next(s['avg_route'] for s in SW if s['s_min'] == 15)], 0.51, 'D'),
    (96, 'ОРШ', [6, 6, 6, tD['orsh']], 0, 'A'),
    (97, 'РОР', [0, tB['n_pop'], tC['n_pop'], 0], 0, None),
    (98, 'сплиттеры 1:64', [mtA['splitters'], 0, 0, mtD['splitters']], 0, None),
    (99, 'сплиттеры 1:8', [0, mtB['splitters1'] + mtB['splitters2'], 0, 0], 0, None),
    (100, 'сплиттеры 1:4', [0, 0, mtC['splitters1'], 0], 0, None),
    (101, 'сплиттеры 1:16', [0, 0, mtC['splitters2'], 0], 0, None),
    (102, 'муфты', [mtA['mufty'], mtB['mufty'], mtC['mufty'], mtD['mufty']], 0, 'C'),
    (103, 'сварки', [mtA['splices'], mtB['splices'], mtC['splices'], mtD['splices']], 0, 'C'),
    (104, 'коннекторы', [mtA['fast_conn'], mtB['fast_conn'], mtC['fast_conn'], mtD['fast_conn']], 0, 'A'),
    (105, 'порты ОРШ', [sum(v['orsh_ports'] for v in ABOOK['villages']),
                        sum(v['orsh_ports'] for v in BBOOK['villages']),
                        sum(v['orsh_ports'] for v in CBOOK['villages']), tD['orsh_ports']], 0, 'C'),
    (106, 'порты OLT', [mtA['olt_ports'], mtB['olt_ports'], mtC['olt_ports'], mtD['olt_ports']], 0, 'A'),
    (107, 'подвесы', [mtA['suspend_kits'], susB, susC, mtD['suspend_kits']], 0, 'D'),
    (108, 'доп. шкафы', [0, 0, 0, tD['orsh'] - 6], 0, 'A'),
]
for r, name, vals, tol, best in G_ROWS:
    for j, val in enumerate(vals):
        if val is None:
            continue
        chk(f'Г: {name} [{"ABCD"[j]}]', ws.cell(row=r, column=4 + j).value, val, tol)
    if best:
        chk(f'Г: {name} лучшая', ws.cell(row=r, column=8).value, best)
# формулы: волокно-км на ДХ (87), ВСЕГО (92)
fk = [4375.1, tB['fiber_km'], tC['fiber_km'], tD['fiber_km']]
for j, L in enumerate('DEFG'):
    chkf(f'Г: волокно-км на ДХ [{L}]', f'{L}87', round(fk[j] / 2334, 4), 0.005)
    chkf(f'Г: ВСЕГО км [{L}]', f'{L}92',
         [cab_tot(mtA), cab_tot(mtB), cab_tot(mtC), cab_tot(mtD)][j]
         + [mtA['drop_cable_km'], mtB['drop_cable_km'], mtC['drop_cable_km'],
            mtD['drop_cable_km']][j], 0.11)
# «лучшая» для строк с формулами
chk('Г: волокно-км на ДХ лучшая', ws.cell(row=87, column=8).value, 'C')
chk('Г: ВСЕГО лучшая', ws.cell(row=92, column=8).value, 'D')

# ---------- блок Д1: индекс стоимости (113-119) ----------
ABC = de28.scheme_abc_costs(0.5)
DC = DEXP['d_costs']['15']
COSTD = dict(A=ABC['A_centr'], B=ABC['B_cascade88'], C=ABC['C_cascade416'], D=DC['cost'])
comp = {}
for tag, c in COSTD.items():
    comp[tag] = [round(c['cable'], 1), round(c['drop'], 1), round(sum(c['hw'].values()), 1)]
for j, tag in enumerate('ABCD'):
    col = chr(68 + j)
    chk(f'Д1: кабель {tag}', ws.cell(row=113, column=4 + j).value, comp[tag][0])
    chk(f'Д1: дропы {tag}', ws.cell(row=114, column=4 + j).value, comp[tag][1])
    chk(f'Д1: обор. {tag}', ws.cell(row=115, column=4 + j).value, comp[tag][2])
    chkf(f'Д1: ИТОГО {tag}', f'{col}116', round(sum(comp[tag]), 1))
box = [0, 0, 0, round(DC['extra_orsh_cost'], 1)]
for j, tag in enumerate('ABCD'):
    col = chr(68 + j)
    chk(f'Д1: доп. шкафы {tag}', ws.cell(row=117, column=4 + j).value, box[j])
    chkf(f'Д1: ИТОГО с шкафами {tag}', f'{col}118', round(sum(comp[tag]) + box[j], 1))
best_all = min(round(sum(comp[t]) + box[j]) for j, t in enumerate('ABCD'))
for j, tag in enumerate('ABCD'):
    col = chr(68 + j)
    chkf(f'Д1: Δ к лучшей {tag}', f'{col}119',
         (round(sum(comp[tag]) + box[j])) / best_all - 1, 0.002)
chk('Д1: лучшая по стоимости (Δ=0 у C)', ws.cell(row=119, column=6).value, 0, 0.001)

# ---------- блок Д2: чувствительность (123-127) ----------
mtD15 = DEXP['d_materials']['15']
orsh_ports15 = next(s['orsh_ports'] for s in SW if s['s_min'] == 15)
for i, k in enumerate((0.0, 0.3, 0.5, 0.7, 1.0)):
    r = 123 + i
    abck = de28.scheme_abc_costs(k)
    cdk = de28.cost_of(mtD15, orsh_ports15, k)
    vals = [round(abck['A_centr']['total'], 1), round(abck['B_cascade88']['total'], 1),
            round(abck['C_cascade416']['total'], 1),
            round(cdk['total'] + DC['extra_orsh_cost'], 1)]
    chk(f'Д2: k={k}', ws.cell(row=r, column=3).value, k, 0)
    for j, v in enumerate(vals):
        chk(f'Д2: k={k} [{"ABCD"[j]}]', ws.cell(row=r, column=4 + j).value, v, 0.06)
    best = 'ABCD'[vals.index(min(vals))]
    chk(f'Д2: k={k} лучшая', ws.cell(row=r, column=8).value, best)

# ---------- Методика: секция 8, пункты 31-36 ----------
ws4 = wb['Методика']
r8 = None
for row in range(5, ws4.max_row + 1):
    v = ws4.cell(row=row, column=3).value
    if isinstance(v, str) and v.startswith('8. Децентрализованная схема'):
        r8 = row
        break
chk('Методика: секция 8 есть', r8 is not None, True)
if r8:
    nums = [ws4.cell(row=r8 + 1 + i, column=2).value for i in range(6)]
    chk('Методика: пункты 31-36', nums, [31, 32, 33, 34, 35, 36])
    names = [ws4.cell(row=r8 + 1 + i, column=3).value for i in range(6)]
    chk('Методика: названия пунктов', names,
        ['Принцип', 'Разбиение на зоны', 'Волокна участков', 'Оборудование зон',
         'Нормы', 'Результат'])

# ---------- неизменность остальных листов ----------
ws2 = wb['Параметры сетей']
chk('«Параметры сетей» ИТОГО волокно-км (не изменён)', ws2.cell(row=11, column=13).value, 889.2)
chk('«Параметры сетей» ИТОГО ДХ (не изменён)', ws2.cell(row=11, column=6).value, 2334, 0)

vals3 = [c.value for row in wb['Сравнение схем'].iter_rows() for c in row
         if isinstance(c.value, (int, float))]
for marker in (4375.1, 1256.1, 889.2):
    chk(f'«Сравнение схем» содержит {marker} (не изменён)',
        any(abs(v - marker) < 0.05 for v in vals3), True)

chk('«Проверка волокно-км» ИТОГО книга (не изменён)',
    wb['Проверка волокно-км']['E13'].value, 4375.1)
chk('«Проверка волокно-км» ИТОГО пересчёт (не изменён)',
    wb['Проверка волокно-км']['F13'].value, 4375.1)

vals1 = [c.value for row in wb['Сводная'].iter_rows() for c in row
         if isinstance(c.value, (int, float))]
for marker, lab in [(2334, 'ДХ'), (323, 'РОР'), (244.0, 'дроп-кабель')]:
    chk(f'«Сводная» содержит {marker} ({lab}, не изменён)',
        any(abs(v - marker) < 0.05 for v in vals1), True)

# ---------- итог ----------
print(f'Проверено значений: {checked}')
if fails:
    print(f'ОШИБОК: {len(fails)}')
    for f in fails[:40]:
        print(' -', f)
    raise SystemExit(1)
print('ВСЕ ПРОВЕРКИ ПРОЙДЕНЫ — лист «Децентрализация ОРШ» соответствует расчётам')
