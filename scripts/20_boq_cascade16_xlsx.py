# -*- coding: utf-8 -*-
"""
Шаг 20. Сводная таблица материалов FTTH по 6 СНП ВКО — КАСКАДНАЯ СХЕМА 1:4 x 1:16 -> Excel.
Листы: «Сводная» (ВОР), «Параметры сетей», «Сравнение схем» (3 схемы), «Методика».
Дизайн — дизайн-система навыка xlsx (templates/base.py), палитра professional, Calibri.
"""
import sys, os, json, math

SKILL = '/home/z/my-project/skills/xlsx'
for sub in [SKILL, os.path.join(SKILL, 'templates')]:
    if sub not in sys.path:
        sys.path.insert(0, sub)
import base

base.FONT_NAME = 'Calibri'
base.HEADER_BOLD = True

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE_DIR = '/home/z/my-project'
OUT = f'{BASE_DIR}/download/Сводная_таблица_материалов_FTTH_ВКО_каскад.xlsx'
DATE = '18.09.2026'

D = json.load(open(f'{BASE_DIR}/work/boq_cascade16_data.json', encoding='utf-8'))   # каскад 1:4 x 1:16
D88 = json.load(open(f'{BASE_DIR}/work/boq_cascade_data.json', encoding='utf-8'))   # каскад 1:8 x 1:8
C = json.load(open(f'{BASE_DIR}/work/boq_data.json', encoding='utf-8'))             # централизованная
V = D['villages']
VN = [v['name'] for v in V]
T = D['totals']

# ---------------------------------------------------------------- стили ----
F_INT = '#,##0'
F_KM = '#,##0.0'
F_KM2 = '#,##0.00'
F_PCT = '+0.0%;-0.0%;0.0%'

def sect_font():
    return Font(name=base.FONT_NAME, size=11, bold=base.HEADER_BOLD, color=base.PRIMARY)

def sect_fill():
    return PatternFill('solid', fgColor=base.SECONDARY)

def cap_font():
    return Font(name=base.FONT_NAME, size=9, color=base.NEUTRAL_600)

def write_headers(ws, headers, row=4):
    for i, h in enumerate(headers, start=2):
        ws.cell(row=row, column=i, value=h)
    base.style_header_row(ws, row_num=row, col_start=2, col_end=1 + len(headers))

def set_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

def nz(x):
    return None if x in (0, 0.0) else x

# ============================================================ ЛИСТ 1 ========
wb = Workbook()
ws = wb.active
ws.title = 'Сводная'

HDRS = ['№', 'Наименование', 'Ед. изм.', 'Расчёт / норма'] + VN + ['ИТОГО']
LAST = 1 + len(HDRS)          # = 12 (L)
base.setup_sheet(ws, title='Сводная таблица материалов — FTTH (GPON), каскадная схема 1:4 + 1:16, 6 СНП Восточно-Казахстанской области', last_col=LAST)
ws['B3'] = ('Каскадное деление 1:4 (ОРШ) × 1:16 (РОР) = 1:64 • сети — по спутниковым снимкам с привязкой к дорогам OSM, топология не менялась • '
            'допущения — лист «Методика» • сравнение схем — лист «Сравнение схем» • подготовлено ' + DATE)
ws['B3'].font = cap_font()
ws.row_dimensions[3].height = 14

write_headers(ws, HDRS)
set_widths(ws, {'A': 3, 'B': 5, 'C': 46, 'D': 9, 'E': 30,
                'F': 13, 'G': 12.5, 'H': 12.5, 'I': 12.5, 'J': 12.5, 'K': 12.5, 'L': 12.5})

SECTIONS = [
    ('А. Оборудование и узлы', [
        ('Шкаф оптический распределительный (ОРШ)', 'шт', '1 на СНП (здание по проекту)', 'orsh', 'int'),
        ('Сплиттер PLC 1×4 — 1-я ступень (в ОРШ)', 'шт', 'сплиттеры 2-й ступ. / 4, окр. вверх', 'splitters1', 'int'),
        ('Сплиттер PLC 1×16 — 2-я ступень (в РОР)', 'шт', '1 на РОР (кластер до 16 ДХ)', 'splitters2', 'int'),
        ('РОР — бокс оптический распределительный (под сплиттер)', 'шт', '1 на кластер ДХ (радиус 150 м)', 'por_boxes', 'int'),
        ('Пигтейль SC/UPC для кросса ОРШ', 'шт', '1,1 × волокна, уходящие от ОРШ', 'pigtails', 'int'),
        ('Порт PON OLT — справочно (активное оборудование)', 'шт', 'сплиттеры 1-й ступени', 'olt_ports', 'int'),
    ]),
    ('Б. Кабельная продукция (длины с запасом на монтаж 10 %)', [
        ('Кабель оптический самонесущий, 8 волокон', 'км', 'волокна = 1,25 × сплиттеры 2-й ступ. потока', 'cable_8', 'km'),
        ('Кабель оптический самонесущий, 12 волокон', 'км', 'то же', 'cable_12', 'km'),
        ('Кабель оптический самонесущий, 16 волокон', 'км', 'то же', 'cable_16', 'km'),
        ('Кабель оптический самонесущий, 24 волокна', 'км', 'то же', 'cable_24', 'km'),
        ('Кабель оптический самонесущий, 32 волокна', 'км', 'то же', 'cable_32', 'km'),
        ('Кабель оптический самонесущий, 48 волокон', 'км', 'то же', 'cable_48', 'km'),
        ('Кабель оптический самонесущий, 64 волокна', 'км', 'то же', 'cable_64', 'km'),
        ('Кабель оптический самонесущий, 72 волокна', 'км', 'то же', 'cable_72', 'km'),
        ('Кабель оптический самонесущий, 96 волокон', 'км', 'то же', 'cable_96', 'km'),
        ('Кабель дроп-кабельный самонесущий, 2 волокна (1 раб. + 1 рез.)', 'км',
         '(дропы + подводки РОР→ДХ) × 1,05', 'drop_cable_km', 'km'),
        ('Справочно: суммарная ёмкость магистрального кабеля', 'волокно-км', 'Σ (длина участка × волокна)', 'fiber_km', 'km'),
    ]),
    ('В. Пассивные узлы', [
        ('Муфта оптическая (ветвления магистрали)', 'шт', 'узлы ветвления ствола; транзит — неразрезным кабелем', 'mufty', 'int'),
    ]),
    ('Г. Материалы для монтажа', [
        ('Бокс абонентский оптический с адаптером SC/UPC', 'шт', '1 на ДХ', 'abonent_boxes', 'int'),
        ('Коннектор оптический механический SC/UPC', 'шт', '2 × ДХ × 1,1 (РОР + абонентский бокс)', 'fast_conn', 'int'),
        ('Сварное соединение (оценка объёма работ)', 'шт', 'пигтейли ОРШ + входы сплиттеров + ответвления, × 1,1', 'splices', 'int'),
        ('Гильза КДЗС, 60 мм', 'шт', 'по числу сварных соединений', 'kdzs', 'int'),
        ('Комплект подвеса магистрального кабеля (кронштейн + спиральный зажим)', 'компл', '30 на 1 км кабеля', 'suspend_kits', 'int'),
        ('Комплект подвеса дроп-кабеля на участках РОР→ДХ', 'компл', '30 на 1 км подводок', 'drop_ext_suspends', 'int'),
        ('Анкерный зажим дроп-кабеля', 'шт', '2 на ДХ', 'drop_anchors', 'int'),
        ('Крепёж дроп-кабеля (скобы / хомуты)', 'шт', '6 на ДХ', 'drop_fix', 'int'),
    ]),
]

row = 5
num = 0
for sect_title, items in SECTIONS:
    ws.cell(row=row, column=3, value=sect_title)
    for c in range(2, LAST + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = sect_fill()
        cell.font = sect_font()
        cell.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[row].height = 24
    row += 1
    for idx, (name, unit, norma, mkey, kind) in enumerate(items):
        num += 1
        ws.cell(row=row, column=2, value=num)
        ws.cell(row=row, column=3, value=name)
        ws.cell(row=row, column=4, value=unit)
        ws.cell(row=row, column=5, value=norma)
        for vi, v in enumerate(V):
            m = dict(v['materials'])
            m['fiber_km'] = v['fiber_km']
            val = m.get(mkey)
            if val in (0, 0.0):
                val = None
            ws.cell(row=row, column=6 + vi, value=val)
        ws.cell(row=row, column=LAST, value=f'=SUM(F{row}:K{row})')
        fill = base.fill_data_row(idx)
        for c in range(2, LAST + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = fill
            cell.font = base.font_body()
            if c == 2:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif c in (3, 5):
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            elif c == 4:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = F_KM if kind == 'km' else F_INT
        ws.cell(row=row, column=LAST).font = base.font_subheader()
        ws.row_dimensions[row].height = 24
        row += 1

row += 1
notes = [
    'Значения в графах сел — потребность по проекту; графа «ИТОГО» — сумма по шести СНП. Единицы измерения разнородны, итог по столбцам не приводится.',
    'Каскадная схема: сплиттеры 1-й ступени 1:4 — в ОРШ, 2-й ступени 1:16 — в РОР (суммарное деление 1:64). Нормы и допущения — лист «Методика»; параметры сетей — лист «Параметры сетей».',
]
for t in notes:
    ws.cell(row=row, column=2, value=t).font = cap_font()
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=LAST)
    ws.row_dimensions[row].height = 14
    row += 1

ws.freeze_panes = 'F5'
ws.page_setup.orientation = 'landscape'
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.print_title_rows = '4:4'

# ============================================================ ЛИСТ 2 ========
ws2 = wb.create_sheet('Параметры сетей')
H2 = ['№', 'СНП', 'Район / с.о.', 'ДХ (заказ)', 'ДХ (проект)', 'Δ к заказу',
      'РОР', 'Сплиттеры 1:16, 2-я ступ.', 'Заполнение портов, %', 'Муфты (ветв.)',
      'Магистраль, км', 'Волокно-км', 'Верхний участок, волокон',
      'Дропы проект, км', 'Подводки РОР→ДХ, км', 'Дроп-кабель, км',
      'Ср. дроп, м', 'Макс. дроп, м', 'ОРШ, портов', 'Сплиттеры 1:4, 1-я ступ.', 'Примечание']
L2 = 1 + len(H2)   # = 22 (V)
base.setup_sheet(ws2, title='Параметры FTTH-сетей по СНП — каскадная схема (1:4 в ОРШ + 1:16 в РОР)', last_col=L2)
ws2['B3'] = ('ДХ — домохозяйства; РОР — пункт распределительный оптический (бокс со сплиттером 1:16, кластер до 16 ДХ в радиусе 150 м). '
             '«Верхний участок» — максимум волокон магистрали на одном участке (у ОРШ). Подготовлено ' + DATE)
ws2['B3'].font = cap_font()
ws2.row_dimensions[3].height = 14
write_headers(ws2, H2)
set_widths(ws2, {'A': 3, 'B': 5, 'C': 16, 'D': 21, 'E': 9, 'F': 9, 'G': 8, 'H': 7, 'I': 10,
                 'J': 10, 'K': 9, 'L': 10, 'M': 9, 'N': 10, 'O': 10, 'P': 10, 'Q': 9, 'R': 9,
                 'S': 9, 'T': 9, 'U': 9, 'V': 40})

r = 5
for i, v in enumerate(V):
    so = f"{v['raion']} / {v['so']}" if v['so'] != '—' else v['raion']
    note = f"ОРШ: {v['orsh_bld']}. {v['note']}"
    vals = [v['num'], v['name'], so, v['dhx_excel'], v['dhx_served'],
            f'=IFERROR(F{r}/E{r}-1,"")', v['n_pop'], v['n2'], v['fill_pct'], v['mufty'],
            v['trunk_km'], v['fiber_km'], v['top_fibers'], v['drop_km'], v['extra_km'],
            v['drop_cable_km'], v['avg_drop_m'], v['max_drop_m'], v['orsh_ports'],
            v['splitters1'], note]
    for ci, val in enumerate(vals, start=2):
        ws2.cell(row=r, column=ci, value=val)
    fill = base.fill_data_row(i)
    for ci in range(2, L2 + 1):
        cell = ws2.cell(row=r, column=ci)
        cell.fill = fill
        cell.font = base.font_body()
        if ci in (2, 4, 7):
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif ci in (3, 22):
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal='right', vertical='center')
    for ci, fmt in [(5, F_INT), (6, F_INT), (7, F_PCT), (8, F_INT), (9, F_INT), (10, '0.0'),
                    (11, F_INT), (12, F_KM2), (13, F_KM), (14, F_INT), (15, F_KM2), (16, F_KM2),
                    (17, F_KM), (18, F_INT), (19, F_INT), (20, F_INT), (21, F_INT)]:
        ws2.cell(row=r, column=ci).number_format = fmt
    ws2.row_dimensions[r].height = 30
    r += 1

tr = r
ws2.cell(row=tr, column=3, value='ИТОГО / справочно')
for ci, f in [(5, '=SUM(E5:E10)'), (6, '=SUM(F5:F10)'), (7, f'=IFERROR(F{tr}/E{tr}-1,"")'),
              (8, '=SUM(H5:H10)'), (9, '=SUM(I5:I10)'),
              (10, f'=IFERROR(F{tr}/(I{tr}*16)*100,"")'),
              (11, '=SUM(K5:K10)'), (12, '=SUM(L5:L10)'), (13, '=SUM(M5:M10)'),
              (15, '=SUM(O5:O10)'), (16, '=SUM(P5:P10)'), (17, '=SUM(Q5:Q10)'),
              (18, f'=IFERROR((O{tr}+P{tr})*1000/F{tr},"")'), (19, '=MAX(S5:S10)'),
              (21, '=SUM(U5:U10)')]:
    ws2.cell(row=tr, column=ci, value=f)
base.style_total_row(ws2, row_num=tr, col_start=2, col_end=L2)
for ci, fmt in [(5, F_INT), (6, F_INT), (7, F_PCT), (8, F_INT), (9, F_INT), (10, '0.0'),
                (11, F_INT), (12, F_KM2), (13, F_KM), (15, F_KM2), (16, F_KM2), (17, F_KM),
                (18, F_INT), (19, F_INT), (21, F_INT)]:
    c = ws2.cell(row=tr, column=ci)
    c.number_format = fmt
    c.alignment = Alignment(horizontal='right', vertical='center')

ws2.freeze_panes = 'F5'
ws2.page_setup.orientation = 'landscape'
ws2.page_setup.fitToWidth = 1
ws2.page_setup.fitToHeight = 0

# ============================================================ ЛИСТ 3 ========
ws3 = wb.create_sheet('Сравнение схем')
H3 = ['№', 'Показатель', 'Ед. изм.', 'Централизованная (1:64 в ОРШ)', 'Каскадная 1:8 + 1:8',
      'Каскадная 1:4 + 1:16 (принято)', 'Δ принятой к 1:8 + 1:8']
L3 = 1 + len(H3)   # = 8 (H)
base.setup_sheet(ws3, title='Сравнение схем деления: централизованная и каскадные (1:8 + 1:8 и 1:4 + 1:16)', last_col=L3)
ws3['B3'] = 'Одни и те же спроектированные сети (трассы, точки подключения ДХ, здания ОРШ). Подготовлено ' + DATE
ws3['B3'].font = cap_font()
ws3.row_dimensions[3].height = 14
write_headers(ws3, H3)
set_widths(ws3, {'A': 3, 'B': 5, 'C': 44, 'D': 9, 'E': 15, 'F': 14, 'G': 15, 'H': 12})

cm, m88, m16 = C['materials_total'], D88['materials_total'], D['materials_total']
SIZES = [8, 12, 16, 24, 32, 48, 64, 72, 96]
cent_cable_km = round(sum(cm.get(f'cable_{s}', 0.0) for s in SIZES), 1)
c88_cable_km = round(sum(m88.get(f'cable_{s}', 0.0) for s in SIZES), 1)
c16_cable_km = round(sum(m16.get(f'cable_{s}', 0.0) for s in SIZES), 1)
cent_fiber = round(sum(v['fiber_km'] for v in C['villages']), 1)
cent_orsh_ports = sum(v['orsh_ports'] for v in C['villages'])
p88_orsh_ports = sum(v['orsh_ports'] for v in D88['villages'])
p16_orsh_ports = sum(v['orsh_ports'] for v in V)

COMP_ROWS = [
    ('S', 'А. Узлы и оборудование', None, None, None, None),
    ('Сплиттер PLC 1×64 (в ОРШ)', 'шт', cm.get('splitters'), None, None, 'int'),
    ('Сплиттер PLC 1×8, 1-я ступень (в ОРШ)', 'шт', None, m88.get('splitters1'), None, 'int'),
    ('Сплиттер PLC 1×8, 2-я ступень (в РОР)', 'шт', None, m88.get('splitters2'), None, 'int'),
    ('Сплиттер PLC 1×4, 1-я ступень (в ОРШ)', 'шт', None, None, m16.get('splitters1'), 'int'),
    ('Сплиттер PLC 1×16, 2-я ступень (в РОР)', 'шт', None, None, m16.get('splitters2'), 'int'),
    ('Сплиттеры всего (обе ступени)', 'шт',
     cm.get('splitters'),
     round(m88.get('splitters1', 0) + m88.get('splitters2', 0)),
     round(m16.get('splitters1', 0) + m16.get('splitters2', 0)), 'int'),
    ('РОР — боксы распределительные', 'шт', None, m88.get('por_boxes'), m16.get('por_boxes'), 'int'),
    ('Муфты оптические', 'шт', cm.get('mufty'), m88.get('mufty'), m16.get('mufty'), 'int'),
    ('Пассивные узлы всего (муфты + РОР)', 'шт', cm.get('mufty'),
     round(m88.get('mufty', 0) + m88.get('por_boxes', 0)),
     round(m16.get('mufty', 0) + m16.get('por_boxes', 0)), 'int'),
    ('Порты PON OLT — справочно', 'шт', cm.get('olt_ports'), m88.get('olt_ports'), m16.get('olt_ports'), 'int'),
    ('Порты ОРШ, суммарно', 'шт', cent_orsh_ports, p88_orsh_ports, p16_orsh_ports, 'int'),
    ('Пигтейли SC/UPC для кросса ОРШ', 'шт', cm.get('pigtails'), m88.get('pigtails'), m16.get('pigtails'), 'int'),
    ('S', 'Б. Кабельная продукция (с запасом на монтаж)', None, None, None, None),
]
for s in SIZES:
    COMP_ROWS.append((f'Кабель самонесущий, {s} волокон', 'км',
                      nz(cm.get(f'cable_{s}', 0.0)), nz(m88.get(f'cable_{s}', 0.0)),
                      nz(m16.get(f'cable_{s}', 0.0)), 'km'))
COMP_ROWS += [
    ('Магистральный кабель — всего', 'км', cent_cable_km, c88_cable_km, c16_cable_km, 'km'),
    ('Суммарная ёмкость магистрали', 'волокно-км', cent_fiber, D88['totals']['fiber_km'], D['totals']['fiber_km'], 'int'),
    ('Дроп-кабель 2-волоконный (с запасом 5 %)', 'км', cm.get('drop_cable_km'), m88.get('drop_cable_km'), m16.get('drop_cable_km'), 'km'),
    ('S', 'В. Монтаж и расходники', None, None, None, None),
    ('Сварные соединения', 'шт', cm.get('splices'), m88.get('splices'), m16.get('splices'), 'int'),
    ('Коннекторы механические SC/UPC', 'шт', cm.get('fast_conn'), m88.get('fast_conn'), m16.get('fast_conn'), 'int'),
    ('Гильзы КДЗС, 60 мм', 'шт', cm.get('kdzs'), m88.get('kdzs'), m16.get('kdzs'), 'int'),
    ('Комплекты подвеса магистрального кабеля', 'компл', cm.get('suspend_kits'), m88.get('suspend_kits'), m16.get('suspend_kits'), 'int'),
    ('Комплекты подвеса дроп-подводок РОР→ДХ', 'компл', None, m88.get('drop_ext_suspends'), m16.get('drop_ext_suspends'), 'int'),
    ('Анкерные зажимы дроп-кабеля', 'шт', cm.get('drop_anchors'), m88.get('drop_anchors'), m16.get('drop_anchors'), 'int'),
    ('Боксы абонентские оптические', 'шт', cm.get('abonent_boxes'), m88.get('abonent_boxes'), m16.get('abonent_boxes'), 'int'),
]

r = 5
n3 = 0
for item in COMP_ROWS:
    if item[0] == 'S':
        ws3.cell(row=r, column=3, value=item[1])
        for c in range(2, L3 + 1):
            cell = ws3.cell(row=r, column=c)
            cell.fill = sect_fill()
            cell.font = sect_font()
            cell.alignment = Alignment(horizontal='left', vertical='center')
        ws3.row_dimensions[r].height = 22
    else:
        n3 += 1
        name, unit, cv, v88, v16, kind = item
        ws3.cell(row=r, column=2, value=n3)
        ws3.cell(row=r, column=3, value=name)
        ws3.cell(row=r, column=4, value=unit)
        ws3.cell(row=r, column=5, value=cv)
        ws3.cell(row=r, column=6, value=v88)
        ws3.cell(row=r, column=7, value=v16)
        ws3.cell(row=r, column=8, value=f'=IF(AND(ISNUMBER(F{r}),ISNUMBER(G{r}),F{r}<>0),G{r}/F{r}-1,"—")')
        fill = base.fill_data_row(n3)
        for c in range(2, L3 + 1):
            cell = ws3.cell(row=r, column=c)
            cell.fill = fill
            cell.font = base.font_body()
            if c == 2:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif c == 3:
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            elif c == 4:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='right', vertical='center')
        for c in (5, 6, 7):
            ws3.cell(row=r, column=c).number_format = F_KM if kind == 'km' else F_INT
        ws3.cell(row=r, column=8).number_format = F_PCT
        ws3.row_dimensions[r].height = 22
    r += 1

r += 1
for t in ['Δ — отношение принятого каскада 1:4 + 1:16 к прежнему каскаду 1:8 + 1:8 («—» — позиция отсутствует в одной из схем). Суммарный коэффициент деления одинаков во всех трёх схемах: 1:64.',
          'Смена соотношения 1:8 + 1:8 → 1:4 + 1:16: РОР и сплиттеров 2-й ступени меньше на 37 %, сварок меньше на 40 %, волокно-км меньше на 29 %, порты ОРШ легче на 14 %; дроп-кабеля больше на 43 % (подводки РОР→ДХ при радиусе кластера 150 м), портов OLT больше на 26 %.']:
    ws3.cell(row=r, column=2, value=t).font = cap_font()
    ws3.merge_cells(start_row=r, start_column=2, end_row=r, end_column=L3)
    ws3.row_dimensions[r].height = 14
    r += 1

ws3.freeze_panes = 'E5'
ws3.page_setup.orientation = 'landscape'
ws3.page_setup.fitToWidth = 1
ws3.page_setup.fitToHeight = 0

# ============================================================ ЛИСТ 4 ========
ws4 = wb.create_sheet('Методика')
H4 = ['№', 'Положение', 'Принято / описание']
L4 = 1 + len(H4)
base.setup_sheet(ws4, title='Методика расчёта и принятые допущения — каскадная схема 1:4 + 1:16', last_col=L4)
ws4['B3'] = 'Расчёт выполнен по геометрии спроектированных сетей (ОРШ → муфты → ДХ) без изменения топологии: трассы, точки подключения ДХ и здания ОРШ — без изменений. Подготовлено ' + DATE
ws4['B3'].font = cap_font()
ws4.row_dimensions[3].height = 14
write_headers(ws4, H4)
set_widths(ws4, {'A': 3, 'B': 5, 'C': 34, 'D': 92})

ROWS4 = [
    ('S', '1. Исходные данные', None),
    ('Геометрия сетей', 'Проектные FTTH-сети «ОРШ → муфты → ДХ»: построены по спутниковым снимкам высокого разрешения (~0,38 м/px) с привязкой к дорожной сети OSM; дропы проложены вдоль улиц, дворов и фасадов, все ветвления — в муфтах. При пересчёте соотношения сплиттеров топология НЕ менялась.'),
    ('Состав сетей', 'Верхнеберезовка — полная сеть села (подтверждена заказчиком, кадр не предоставлялся); Винное — полная сеть на кадре пользователя (550 ДХ, из них 5 за верхней кромкой кадра); Солнечное, Перевальное, Пригородное, Алтайский — в границах предоставленных заказчиком кадров.'),
    ('Число ДХ по данным заказа', '«СНП ВКО.xlsx», лист «Лоты 12»: 3 216 ДХ по шести СНП (Верхнеберезовка 940, Солнечное 366, Перевальное 339, Винное 490, Пригородное 365, Алтайский 716). Обслужено проектом — 2 334 ДХ (лист «Параметры сетей»).'),
    ('Длины', 'Длины участков и дропов — по геометрии проекта с пересчётом пикселей в метры по масштабу мозаик (~0,38 м/px).'),
    ('S', '2. Архитектура (каскадная схема, соотношение 1:4 + 1:16)', None),
    ('Схема деления', 'GPON, двухступенчатое каскадное деление: сплиттер PLC 1×4 в ОРШ (1-я ступень) × сплиттер PLC 1×16 в РОР (2-я ступень) = 1:64. Суммарный коэффициент деления сохранён таким же, как в централизованной схеме (1:64) и в прежнем варианте каскада (1:8 × 1:8); оптический бюджет — класс B+.'),
    ('РОР', 'Пункт распределительный оптический (бокс со сплиттером 1:16) размещается в точке магистрали (проектной муфте) и обслуживает кластер до 16 ДХ в радиусе 150 м по трассе сети. Всего 323 РОР на 6 СНП (в варианте 1:8 + 1:8 — 515); среднее 7,2 ДХ на РОР; среднее заполнение портов 45 % (37–51 % по сёлам) — типично для сельской плотности застройки (~2 ДХ на муфту по проекту).'),
    ('Дроп-линия', 'От РОР к каждому ДХ — индивидуальный 2-волоконный дроп-кабель (1 рабочее + 1 резервное волокно). Участок «РОР → точка подключения ДХ» прокладывается дроп-кабелем по трассе магистрали: средняя подводка +61 м на ДХ (максимум 150 м), суммарно 141,3 км. Максимальная суммарная дроп-линия — 266 м.'),
    ('Сплиттеры 1-й ступени', 'PLC 1×4 в ОРШ: 83 шт на 6 СНП (по каждому селу — сплиттеры 2-й ступени / 4, округление вверх); каждый питает до 4 РОР.'),
    ('S', '3. Кабельная продукция', None),
    ('Магистральный кабель', 'Самонесущий оптический кабель для воздушной прокладки по существующим опорам; ёмкости 8 / 12 / 16 / 24 / 32 / 48 / 64 / 72 / 96 волокон. При потребности свыше 96 волокон — параллельная прокладка кабелей меньшей ёмкости (верхние ярусы дерева).'),
    ('Число волокон на участке', 'Сплиттеры 2-й ступени ниже по потоку × 1,25 (резерв 25 % на развитие), минимум 8 волокон (минимальная серийная ёмкость). Верхний участок — 128 волокон (Верхнеберезовка, до 2 параллельных кабелей у ОРШ); кабели ёмкостью более 48 волокон требуются только вблизи ОРШ крупных сёл.'),
    ('Запас на монтаж', '+10 % к расчётной длине магистрали и +5 % к дроп-кабелю; округление вверх до 0,1 км.'),
    ('Справочно', 'Суммарная расчётная ёмкость магистрали по шести СНП — 889 волокно-км (в централизованной схеме — 4 375; в каскаде 1:8 + 1:8 — 1 256 волокно-км).'),
    ('S', '4. Узлы и монтажные нормы', None),
    ('Муфты', 'Только узлы ветвления магистрали (100 шт по шести СНП): участки без ветвлений выполняются неразрезным кабелем, транзитные волокна не сращиваются. РОР учитываются отдельной позицией; смена ёмкости кабеля совмещается с РОР и муфтами ветвления.'),
    ('ОРШ', '1 на СНП, на общественном/административном здании в центре села (лист «Параметры сетей»). Расчётная ёмкость портов — 1,1 × волокна, уходящие к РОР, с округлением вверх до ряда 144 / 288 / 576 / 864 / 1152: суммарно 864 порта (у всех шести СНП — по 144) против 1 008 в каскаде 1:8 + 1:8 и 3 600 в централизованной схеме.'),
    ('Сварные соединения', 'Оценка: пигтейли кросса ОРШ + входы сплиттеров РОР + ответвления и смены ёмкости кабеля в узлах, с запасом 10 %. Гильзы КДЗС — по числу сварок. В каскадной схеме дропы к сварке не подключаются (механические коннекторы) — объём сварок в 1,7 раза меньше каскада 1:8 + 1:8 (1 552 против 2 585) и в 3,3 раза меньше централизованной схемы (5 136).'),
    ('Коннекторы', 'Механические SC/UPC: 2 на ДХ (выход сплиттера в РОР + абонентский бокс), запас 10 %. Дроп-подводки РОР→ДХ крепятся по опорам (30 компл./км), вводы в здания — скобами/хомутами (6 точек на ДХ).'),
    ('Подвес магистрали', '30 комплектов (кронштейн + спиральный зажим) на 1 км уложенного кабеля; шаг промежуточных креплений ~35 м. Новые опоры не предусмотрены — существующие опорные конструкции.'),
    ('Крепёж дропа', '2 анкерных зажима (у РОР/муфты и у стены здания) и 6 промежуточных креплений (скобы/хомуты) на одно ДХ.'),
    ('S', '5. Прочее', None),
    ('Активное оборудование', 'Порты PON OLT — 83 (справочно; в каскаде 1:8 + 1:8 — 66, в централизованной схеме — 39): следствие более дробного деления 1-й ступени при заполнении портов сплиттеров менее 100 %. Модель, резервирование и комплектацию уточняет заказчик.'),
    ('Сравнение схем', 'Количественное сравнение трёх схем деления (централизованная 1:64; каскад 1:8 + 1:8; каскад 1:4 + 1:16) по всем позициям — лист «Сравнение схем».'),
    ('За рамками расчёта', 'Строительство или аренда опор, заземление, оборудование абонентов (ONT), СМР, накладные и транспортные расходы.'),
]
r = 5
n4 = 0
for item in ROWS4:
    if item[0] == 'S':
        ws4.cell(row=r, column=3, value=item[1])
        for c in range(2, L4 + 1):
            cell = ws4.cell(row=r, column=c)
            cell.fill = sect_fill()
            cell.font = sect_font()
            cell.alignment = Alignment(horizontal='left', vertical='center')
        ws4.row_dimensions[r].height = 24
    else:
        n4 += 1
        name, desc = item
        ws4.cell(row=r, column=2, value=n4)
        ws4.cell(row=r, column=3, value=name)
        ws4.cell(row=r, column=4, value=desc)
        fill = base.fill_data_row(n4)
        for c in range(2, L4 + 1):
            cell = ws4.cell(row=r, column=c)
            cell.fill = fill
            cell.font = base.font_body()
            if c == 2:
                cell.alignment = Alignment(horizontal='center', vertical='top')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        lines = max(math.ceil(len(desc) / 110), math.ceil(len(name) / 30), 1)
        ws4.row_dimensions[r].height = max(22, lines * 14 + 8)
    r += 1

ws4.page_setup.orientation = 'landscape'
ws4.page_setup.fitToWidth = 1
ws4.page_setup.fitToHeight = 0

wb.properties.creator = 'Z.ai'
wb.save(OUT)
print('Сохранено:', OUT)
print('Листы:', wb.sheetnames)
