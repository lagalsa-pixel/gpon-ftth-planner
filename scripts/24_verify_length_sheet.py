# -*- coding: utf-8 -*-
"""
Шаг 24. Семантическая проверка листа «Оптимум по длине» (и дополнения «Методики»)
против расчётных данных work/length_sweep.json.

Проверяется (data_only, после recalc — кэш формул доступен):
  - блок А (3 схемы + оптимум): дроп/подводки/магистраль/СУММА(J-формула)/волокно-км/
    ср-макс дроп/сварки/заполнение;
  - блок Б: все 17 конфигураций по тем же метрикам + Δ-формулы;
  - блок В: 6 СНП оптимума R50 + итоговая строка (SUM-формулы);
  - блок Д (данные графика): 9 точек x 3 серии;
  - «Методика»: 25 пунктов, наличие секции 6.
"""
import json
from openpyxl import load_workbook

BASE = '/home/z/my-project'
F = f'{BASE}/download/Сводная_таблица_материалов_FTTH_ВКО_каскад.xlsx'
SW = json.load(open(f'{BASE}/work/length_sweep.json', encoding='utf-8'))
REFS, CFGS = SW['references'], SW['configs']
DH = CFGS[0]['dhx_served']

def cfg(r_max, cap):
    return next(c for c in CFGS if c['r_max'] == r_max and c['cap'] == cap)

wb = load_workbook(F, data_only=True)
ws = wb['Оптимум по длине']
errors = []
checked = 0

def chk(name, got, exp, tol=0.051):
    global checked
    checked += 1
    if exp is None:
        return
    if got is None or abs(float(got) - float(exp)) > tol:
        errors.append(f'{name}: в книге {got}, ожидалось {exp}')

def near(name, got, exp, rel=0.01):
    global checked
    checked += 1
    if got is None or abs(float(got) - float(exp)) > abs(float(exp)) * rel + 0.05:
        errors.append(f'{name}: в книге {got}, ожидалось ~{exp}')

# ---------------- блок А: строки 6-9 ----------------
A = [
    (6, REFS['centralized'], None),
    (7, REFS['cascade88'], 8),
    (8, REFS['cascade16'], 16),
    (9, cfg(50, 8), 8),
]
for row, c, split2 in A:
    tag = f'A:{ws.cell(row=row, column=3).value[:30]}'
    chk(f'{tag} дроп', ws.cell(row=row, column=7).value, c['drop_cable_km'])
    chk(f'{tag} подводки', ws.cell(row=row, column=8).value, c.get('extra_km'))
    chk(f'{tag} магистраль', ws.cell(row=row, column=9).value, c['cable_boq_km'])
    chk(f'{tag} СУММА(J)', ws.cell(row=row, column=10).value,
        round(c['drop_cable_km'] + c['cable_boq_km'], 1))
    chk(f'{tag} волокно-км', ws.cell(row=row, column=12).value, c['fiber_km'])
    chk(f'{tag} ср.дроп', ws.cell(row=row, column=13).value, c['avg_drop_m'], tol=0.06)
    chk(f'{tag} макс.дроп', ws.cell(row=row, column=14).value, c['max_drop_m'], tol=0.06)
    if c.get('splices') is not None:
        chk(f'{tag} сварки', ws.cell(row=row, column=15).value, c['splices'], tol=0.1)
    if c.get('n_pop') is not None:
        chk(f'{tag} РОР', ws.cell(row=row, column=5).value, c['n_pop'], tol=0.1)
    if split2 and c.get('n2'):
        chk(f'{tag} заполнение', ws.cell(row=row, column=6).value,
            round(100.0 * DH / (c['n2'] * split2), 1), tol=0.06)

# Δ блока А: база — строка R0 (J12)
L0 = cfg(0, 8)['L_total']
for row, c in [(6, REFS['centralized']), (7, REFS['cascade88']), (8, REFS['cascade16']),
               (9, cfg(50, 8))]:
    near(f'A:r{row} Δ', ws.cell(row=row, column=11).value,
         round(c['drop_cable_km'] + c['cable_boq_km'], 1) / L0 - 1, rel=0.005)

# ---------------- блок Б: строки 12-28 ----------------
r = 12
for r_max in [0, 50, 75, 100, 125, 150, 175, 200, 250]:
    caps = [8] if r_max == 0 else [8, 16]
    for cap in caps:
        c = cfg(r_max, cap)
        tag = f'Б:R{r_max}/c{cap}'
        chk(f'{tag} РОР', ws.cell(row=r, column=5).value, c['n_pop'], tol=0.1)
        chk(f'{tag} заполнение', ws.cell(row=r, column=6).value, c['fill_pct'], tol=0.06)
        chk(f'{tag} дроп', ws.cell(row=r, column=7).value, c['drop_cable_km'])
        chk(f'{tag} подводки', ws.cell(row=r, column=8).value, c['extra_km'])
        chk(f'{tag} магистраль', ws.cell(row=r, column=9).value, c['cable_boq_km'])
        chk(f'{tag} СУММА(J)', ws.cell(row=r, column=10).value,
            round(c['drop_cable_km'] + c['cable_boq_km'], 1))
        chk(f'{tag} волокно-км', ws.cell(row=r, column=12).value, c['fiber_km'])
        chk(f'{tag} ср.дроп', ws.cell(row=r, column=13).value, c['avg_drop_m'], tol=0.06)
        chk(f'{tag} макс.дроп', ws.cell(row=r, column=14).value, c['max_drop_m'], tol=0.06)
        chk(f'{tag} сварки', ws.cell(row=r, column=15).value, c['splices'], tol=0.1)
        near(f'{tag} Δ', ws.cell(row=r, column=11).value,
             round(c['drop_cable_km'] + c['cable_boq_km'], 1) / L0 - 1, rel=0.005)
        r += 1
assert r == 29, f'блок Б закончился на строке {r}, ожидалось 29'

# ---------------- блок В: строки 35-40, итог 41 ----------------
R50 = cfg(50, 8)
for i, v in enumerate(R50['per_village']):
    row = 35 + i
    tag = f'В:{v["name"]}'
    chk(f'{tag} ДХ', ws.cell(row=row, column=4).value, v['dhx_served'], tol=0.1)
    chk(f'{tag} РОР', ws.cell(row=row, column=5).value, v['n_pop'], tol=0.1)
    chk(f'{tag} заполнение', ws.cell(row=row, column=6).value, v['fill_pct'], tol=0.06)
    chk(f'{tag} дроп', ws.cell(row=row, column=7).value, v['drop_cable_km'])
    chk(f'{tag} подводки', ws.cell(row=row, column=8).value, v['extra_km'])
    chk(f'{tag} магистраль', ws.cell(row=row, column=9).value, v['cable_boq_km'])
    chk(f'{tag} СУММА(J)', ws.cell(row=row, column=10).value,
        round(v['drop_cable_km'] + v['cable_boq_km'], 1))
    chk(f'{tag} волокно-км', ws.cell(row=row, column=11).value, v['fiber_km'])

tr = 41
tot = R50
near('В:ИТОГО ДХ', ws.cell(row=tr, column=4).value, DH)
chk('В:ИТОГО РОР', ws.cell(row=tr, column=5).value, tot['n_pop'], tol=0.1)
chk('В:ИТОГО заполнение', ws.cell(row=tr, column=6).value, tot['fill_pct'], tol=0.06)
chk('В:ИТОГО дроп', ws.cell(row=tr, column=7).value, tot['drop_cable_km'])
chk('В:ИТОГО магистраль', ws.cell(row=tr, column=9).value, tot['cable_boq_km'])
chk('В:ИТОГО СУММА', ws.cell(row=tr, column=10).value, tot['L_total'])
chk('В:ИТОГО волокно-км', ws.cell(row=tr, column=11).value, tot['fiber_km'])
near('В:ИТОГО ср.дроп', ws.cell(row=tr, column=12).value, tot['avg_drop_m'], rel=0.01)
chk('В:ИТОГО макс.дроп', ws.cell(row=tr, column=13).value, tot['max_drop_m'], tol=0.06)

# ---------------- блок Д: данные графика, строки 51-59 ----------------
r = 51
for r_max in [0, 50, 75, 100, 125, 150, 175, 200, 250]:
    c8, c16 = cfg(r_max, 8), cfg(r_max, 16)
    chk(f'Д:R{r_max} радиус', ws.cell(row=r, column=2).value, r_max, tol=0.1)
    chk(f'Д:R{r_max} сумма 1:8', ws.cell(row=r, column=3).value, c8['L_total'])
    chk(f'Д:R{r_max} сумма 1:16', ws.cell(row=r, column=4).value, c16['L_total'])
    chk(f'Д:R{r_max} вол-км 1:16', ws.cell(row=r, column=5).value, c16['fiber_km'])
    r += 1

# ---------------- Методика ----------------
ws4 = wb['Методика']
nums = [ws4.cell(row=rr, column=2).value for rr in range(5, ws4.max_row + 1)
        if isinstance(ws4.cell(row=rr, column=2).value, int)]
checked += 3
if nums != list(range(1, 26)):
    errors.append(f'Методика: нумерация пунктов {nums[:5]}...{nums[-5:]} (ожидалось 1..25)')
has6 = any(ws4.cell(row=rr, column=3).value == '6. Оптимизация по длине кабельной продукции'
           for rr in range(5, ws4.max_row + 1))
if not has6:
    errors.append('Методика: не найдена секция 6')
if ws4.cell(row=ws4.max_row, column=3).value != 'Рекомендация':
    errors.append('Методика: последний пункт не «Рекомендация»')

# ---------------- прежние листы не изменились ----------------
ws1 = wb['Сводная']
checked += 3
if ws1.cell(row=4, column=3).value != 'Наименование':
    errors.append('Сводная: заголовки сместились')
if wb.sheetnames != ['Сводная', 'Параметры сетей', 'Сравнение схем', 'Оптимум по длине', 'Методика']:
    errors.append(f'Порядок листов: {wb.sheetnames}')
drop_row = None
for rr in range(5, ws1.max_row + 1):
    v = ws1.cell(row=rr, column=3).value
    if isinstance(v, str) and 'дроп-кабельный' in v.lower():
        drop_row = rr
        break
if drop_row is None:
    errors.append('Сводная: не найдена строка дроп-кабеля')
elif abs((ws1.cell(row=drop_row, column=12).value or 0) - 244.0) > 0.2:
    errors.append(f'Сводная: дроп-кабель ИТОГО (строка {drop_row}) = '
                  f'{ws1.cell(row=drop_row, column=12).value} (ожидалось 244.0)')

print(f'Проверено значений: {checked}')
if errors:
    print(f'ОШИБКИ ({len(errors)}):')
    for e in errors:
        print(' -', e)
    raise SystemExit(1)
print('ВСЕ ЗНАЧЕНИЯ СХОДЯТСЯ ✓')
