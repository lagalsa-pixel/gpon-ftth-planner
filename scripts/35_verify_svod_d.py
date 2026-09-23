# -*- coding: utf-8 -*-
"""
Шаг 35. Семантическая проверка листа «Сводная (схема D)» (после QA-пайплайна).

Сверяет с work/boq_decentral_data.json (шаг 29):
  - все 24 позиции по 6 графам сел (значения);
  - графы ИТОГО (кэш recalc формул SUM);
  - структуру секций и единицы измерения;
  - пункты «Методики» 37-39 (секция 9);
  - порядок листов (сводная D — вторая);
  - неизменность остальных листов: спот-чеки «Сводной» (схема C),
    «Параметров сетей», «Сравнения схем», «Децентрализации ОРШ».
"""
import json
from openpyxl import load_workbook

BASE = '/home/z/my-project'
OUT = f'{BASE}/download/Сводная_таблица_материалов_FTTH_ВКО_каскад.xlsx'

DBOOK = json.load(open(f'{BASE}/work/boq_decentral_data.json', encoding='utf-8'))
CBOOK = json.load(open(f'{BASE}/work/boq_cascade16_data.json', encoding='utf-8'))
V = DBOOK['villages']

wb = load_workbook(OUT, data_only=True)      # кэш recalc
wbf = load_workbook(OUT)                     # формулы
ws, wsf = wb['Сводная (схема D)'], wbf['Сводная (схема D)']

fails = []
checked = 0


def chk(desc, actual, expected, tol=0.051):
    global checked
    checked += 1
    ok = (abs(actual - expected) <= tol) if isinstance(expected, (int, float)) \
        else (actual == expected)
    if not ok:
        fails.append(f'{desc}: лист={actual!r} vs ожидалось={expected!r}')


# ---------- порядок листов ----------
EXPECT = ['Сводная', 'Сводная (схема D)', 'Параметры сетей', 'Сравнение схем',
          'Оптимум по длине', 'Проверка волокно-км', 'Децентрализация ОРШ', 'Методика']
chk('порядок листов', wb.sheetnames, EXPECT)

# ---------- шапка ----------
chk('заголовок листа', ws['B4'].value is not None, True)

# ---------- позиции ----------
# карта: имя позиции -> (ключ материалов или функция, ед.)
POS = [
    ('Шкаф оптический распределительный (ОРШ) центральный — в здании ЦУ', lambda v: 1, 'шт'),
    ('Шкаф зонный ОРШ уличный (под сплиттеры 1:64)', lambda v: v['n_zones'], 'шт'),
    ('Сплиттер PLC 1×64 (в ОРШ: ЦУ и зонных)', 'splitters', 'шт'),
    ('Пигтейль SC/UPC для кросса ОРШ', 'pigtails', 'шт'),
    ('Порт PON OLT — справочно (активное оборудование, единый узел OLT)', 'olt_ports', 'шт'),
    ('Кабель оптический самонесущий, 8 волокон', 'cable_8', 'км'),
    ('Кабель оптический самонесущий, 12 волокон', 'cable_12', 'км'),
    ('Кабель оптический самонесущий, 16 волокон', 'cable_16', 'км'),
    ('Кабель оптический самонесущий, 24 волокна', 'cable_24', 'км'),
    ('Кабель оптический самонесущий, 32 волокна', 'cable_32', 'км'),
    ('Кабель оптический самонесущий, 48 волокон', 'cable_48', 'км'),
    ('Кабель оптический самонесущий, 64 волокна', 'cable_64', 'км'),
    ('Кабель оптический самонесущий, 72 волокна', 'cable_72', 'км'),
    ('Кабель оптический самонесущий, 96 волокон', 'cable_96', 'км'),
    ('Кабель дроп-кабельный самонесущий, 2 волокна (1 раб. + 1 рез.)', 'drop_cable_km', 'км'),
    ('Справочно: суммарная ёмкость магистрального кабеля', 'fiber_km', 'волокно-км'),
    ('Муфта оптическая (ветвления магистрали)', 'mufty', 'шт'),
    ('Бокс абонентский оптический с адаптером SC/UPC', 'abonent_boxes', 'шт'),
    ('Коннектор оптический механический SC/UPC', 'fast_conn', 'шт'),
    ('Сварное соединение (оценка объёма работ)', 'splices', 'шт'),
    ('Гильза КДЗС, 60 мм', 'kdzs', 'шт'),
    ('Комплект подвеса магистрального кабеля (кронштейн + спиральный зажим)', 'suspend_kits', 'компл'),
    ('Анкерный зажим дроп-кабеля', 'drop_anchors', 'шт'),
    ('Крепёж дроп-кабеля (скобы / хомуты)', 'drop_fix', 'шт'),
]

# найти строки позиций по имени в колонке C
row_of = {}
for r in range(5, ws.max_row + 1):
    name = ws.cell(row=r, column=3).value
    if isinstance(name, str) and name in {p[0] for p in POS}:
        row_of[name] = r

chk('позиций найдено', len(row_of), 24)
num_seen = []
for name, key, unit in POS:
    r = row_of.get(name)
    if r is None:
        fails.append(f'нет строки: {name}')
        continue
    num_seen.append(ws.cell(row=r, column=2).value)
    chk(f'[{name}] ед.', ws.cell(row=r, column=4).value, unit)
    for vi, v in enumerate(V):
        m = dict(v['materials'])
        m['fiber_km'] = v['fiber_km']
        exp = key(v) if callable(key) else m.get(key)
        got = ws.cell(row=r, column=6 + vi).value
        if exp in (0, 0.0):
            chk(f'[{name}][{v["name"]}] ноль -> пусто', got, None)
        else:
            chk(f'[{name}][{v["name"]}]', got, exp)
    # ИТОГО по кэшу + формула
    col_tot = ws.cell(row=r, column=12).value
    frm = wsf.cell(row=r, column=12).value
    exp_sum = sum((key(v) if callable(key) else dict(v['materials'], fiber_km=v['fiber_km']).get(key, 0))
                  for v in V)
    exp_sum = round(exp_sum, 1)
    chk(f'[{name}] ИТОГО (кэш)', col_tot, exp_sum if exp_sum else None)
    chk(f'[{name}] ИТОГО формула', frm, f'=SUM(F{r}:K{r})')

chk('нумерация позиций 1..24', num_seen, list(range(1, 25)))

# ---------- секции ----------
sects = [ws.cell(row=r, column=3).value for r in range(5, ws.max_row + 1)
         if isinstance(ws.cell(row=r, column=3).value, str)
         and ws.cell(row=r, column=3).value.startswith(('А.', 'Б.', 'В.', 'Г.'))]
chk('секции', sects, ['А. Оборудование и узлы', 'Б. Кабельная продукция (длины с запасом на монтаж 10 %)',
                      'В. Пассивные узлы', 'Г. Материалы для монтажа'])

# ---------- контрольные итоги ----------
mt = DBOOK['materials_total']
r_spl = row_of['Сплиттер PLC 1×64 (в ОРШ: ЦУ и зонных)']
chk('сплиттеры ИТОГО', ws.cell(row=r_spl, column=12).value, mt['splitters'])
r_mu = row_of['Муфта оптическая (ветвления магистрали)']
chk('муфты ИТОГО', ws.cell(row=r_mu, column=12).value, mt['mufty'])
r_fk = row_of['Справочно: суммарная ёмкость магистрального кабеля']
chk('волокно-км ИТОГО', ws.cell(row=r_fk, column=12).value, 1562.5, tol=0.05)
r_z = row_of['Шкаф зонный ОРШ уличный (под сплиттеры 1:64)']
chk('зонные шкафы ИТОГО', ws.cell(row=r_z, column=12).value, 28)
r_c = row_of['Шкаф оптический распределительный (ОРШ) центральный — в здании ЦУ']
chk('ЦУ ОРШ ИТОГО', ws.cell(row=r_c, column=12).value, 6)
r_sp = row_of['Сварное соединение (оценка объёма работ)']
chk('сварки ИТОГО', ws.cell(row=r_sp, column=12).value, mt['splices'])

# ---------- Методика 37-39 ----------
ws4 = wb['Методика']
m_items = {}
for r in range(5, ws4.max_row + 1):
    n = ws4.cell(row=r, column=2).value
    if isinstance(n, (int, float)) and 37 <= int(n) <= 39:
        m_items[int(n)] = (ws4.cell(row=r, column=3).value, ws4.cell(row=r, column=4).value)
chk('Методика: пункты 37-39', sorted(m_items.keys()), [37, 38, 39])
chk('Методика 37 имя', m_items.get(37, ('', ''))[0], 'Состав позиций')
chk('Методика 38 имя', m_items.get(38, ('', ''))[0], 'Нормы схемы D')
chk('Методика 39 имя', m_items.get(39, ('', ''))[0], 'Карты зон ОРШ')
for n, (nm, desc) in m_items.items():
    chk(f'Методика {n} текст непуст', len(desc or '') > 100, True)

# заголовок секции 9
sec9 = [ws4.cell(row=r, column=3).value for r in range(5, ws4.max_row + 1)
        if isinstance(ws4.cell(row=r, column=3).value, str)
        and ws4.cell(row=r, column=3).value.startswith('9. Сводная материалов')]
chk('Методика: секция 9', len(sec9), 1)

# ---------- неизменность других листов ----------
wsA = wb['Сводная']  # схема C
chk('Сводная C: заголовок', str(wsA['B4'].value or '')[:30], str(wsA['B4'].value or '')[:30])
# спот-чек: сплиттеры 1:16 (поз. 3) ИТОГО = 323
r16 = None
for r in range(5, wsA.max_row + 1):
    if wsA.cell(row=r, column=3).value == 'Сплиттер PLC 1×16 — 2-я ступень (в РОР)':
        r16 = r
        break
chk('Сводная C: сплиттеры 1:16 ИТОГО', wsA.cell(row=r16, column=12).value, CBOOK['totals']['n2'])

wsP = wb['Параметры сетей']
chk('Параметры сетей: волокно-км итог', wsP.cell(row=11, column=13).value, 889.2, tol=0.05)

wsS = wb['Сравнение схем']
found = 0
for r in range(1, wsS.max_row + 1):
    for c in range(1, wsS.max_column + 1):
        val = wsS.cell(row=r, column=c).value
        if isinstance(val, (int, float)) and abs(val - 4375.1) < 0.05:
            found += 1
chk('Сравнение схем: 4375,1 присутствует', found >= 1, True)

wsD = wb['Децентрализация ОРШ']
found = 0
for r in range(1, wsD.max_row + 1):
    for c in range(1, wsD.max_column + 1):
        val = wsD.cell(row=r, column=c).value
        if isinstance(val, (int, float)) and abs(val - 1562.5) < 0.05:
            found += 1
chk('Децентрализация ОРШ: 1562,5 присутствует', found >= 1, True)

# ---------- итог ----------
print(f'Проверено значений: {checked}')
if fails:
    print(f'ОШИБОК: {len(fails)}')
    for f in fails[:25]:
        print(' -', f)
    raise SystemExit(1)
print('ВСЕ ЗНАЧЕНИЯ СХОДЯТСЯ ✓')
