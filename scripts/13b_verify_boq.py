# -*- coding: utf-8 -*-
"""Шаг 13b. Семантическая проверка книги Excel против boq_data.json."""
import json, math
import openpyxl

BASE = '/home/z/my-project'
F = f'{BASE}/download/Сводная_таблица_материалов_FTTH_ВКО.xlsx'
D = json.load(open(f'{BASE}/work/boq_data.json', encoding='utf-8'))
V = D['villages']
MT = D['materials_total']

wb = openpyxl.load_workbook(F, data_only=True)  # только чтение кэша, НЕ сохранять
ws, ws2 = wb['Сводная'], wb['Параметры сетей']

# --- карта строк листа 1 по наименованию ---
rowmap = {}
for r in range(5, ws.max_row + 1):
    name = ws.cell(row=r, column=3).value
    if name and ws.cell(row=r, column=2).value is not None:
        rowmap[(ws.cell(row=r, column=2).value, name)] = r

def total_of(name):
    for (num, nm), r in rowmap.items():
        if nm == name:
            return ws.cell(row=r, column=12).value
    raise KeyError(name)

exp = {
    'Шкаф оптический распределительный (ОРШ)': MT['orsh'],
    'Сплиттер PLC 1×64 (устанавливается в ОРШ)': MT['splitters'],
    'Пигтейль SC/UPC для кросса ОРШ': MT['pigtails'],
    'Порт PON OLT — справочно (активное оборудование)': MT['olt_ports'],
    'Кабель оптический самонесущий, 8 волокон': MT['cable_8'],
    'Кабель оптический самонесущий, 12 волокон': MT['cable_12'],
    'Кабель оптический самонесущий, 16 волокон': MT['cable_16'],
    'Кабель оптический самонесущий, 24 волокна': MT['cable_24'],
    'Кабель оптический самонесущий, 32 волокна': MT['cable_32'],
    'Кабель оптический самонесущий, 48 волокон': MT['cable_48'],
    'Кабель оптический самонесущий, 64 волокна': MT['cable_64'],
    'Кабель оптический самонесущий, 72 волокна': MT['cable_72'],
    'Кабель оптический самонесущий, 96 волокон': MT['cable_96'],
    'Кабель дроп-кабельный самонесущий, 2 волокна (1 раб. + 1 рез.)': MT['drop_cable_km'],
    'Справочно: суммарная ёмкость магистрального кабеля': round(sum(v['fiber_km'] for v in V), 1),
    'Муфта оптическая (проходная / тупиковая)': MT['mufty'],
    'Бокс абонентский оптический с адаптером SC/UPC': MT['abonent_boxes'],
    'Коннектор оптический механический SC/UPC': MT['fast_conn'],
    'Сварное соединение (оценка объёма работ)': MT['splices'],
    'Гильза КДЗС, 60 мм': MT['kdzs'],
    'Комплект подвеса магистрали (кронштейн + спиральный зажим)': MT['suspend_kits'],
    'Анкерный зажим дроп-кабеля': MT['drop_anchors'],
    'Крепёж дроп-кабеля (скобы / хомуты)': MT['drop_fix'],
}

fails = 0
for name, e in exp.items():
    got = total_of(name)
    ok = got is not None and abs(float(got) - float(e)) < 0.051
    if not ok:
        fails += 1
        print(f'FAIL {name}: в книге {got}, ожидалось {e}')
print(f'Лист 1 «Сводная»: проверено позиций {len(exp)}, ошибок {fails}')

# --- построчная сверка сел (спот-чeck 3 позиций x 6 сел) ---
spot = {('Муфта оптическая (проходная / тупиковая)', 8): [v['couplers'] for v in V],
        ('Пигтейль SC/UPC для кросса ОРШ', 6): [v['dhx_served'] for v in V],
        ('Кабель дроп-кабельный самонесущий, 2 волокна (1 раб. + 1 рез.)', 7): [v['drop_cable_km'] for v in V]}
for (name, _), expected in spot.items():
    r = total_of.__self__ if False else None
    for (num, nm), rr in rowmap.items():
        if nm == name:
            r = rr
    vals = [ws.cell(row=r, column=6 + i).value for i in range(6)]
    for i, (g, e) in enumerate(zip(vals, expected)):
        if g is None or abs(float(g) - float(e)) > 0.051:
            fails += 1
            print(f'FAIL {name} / {V[i]["name"]}: {g} != {e}')
print('Спот-проверка по сёлам: ошибок', fails)

# --- лист 2: итоги ---
t = {}
for r in range(5, ws2.max_row + 1):
    if ws2.cell(row=r, column=3).value == 'ИТОГО / справочно':
        t = {c: ws2.cell(row=r, column=c).value for c in range(2, 17)}
        break
e2 = {5: 3216, 6: 2334, 8: 1140, 9: 73.08, 10: 90.89, 14: 39}
for c, e in e2.items():
    g = t.get(c)
    if g is None or abs(float(g) - e) > 0.06:
        fails += 1
        print(f'FAIL Параметры итоги col{c}: {g} != {e}')
g7, g11 = t.get(7), t.get(11)
if g7 is None or abs(g7 - (2334 / 3216 - 1)) > 0.001:
    fails += 1; print(f'FAIL Δ: {g7}')
if g11 is None or abs(g11 - 90.89 * 1000 / 2334) > 0.1:
    fails += 1; print(f'FAIL ср.дроп: {g11}')
print('Лист 2 «Параметры сетей»: итоги сверены')

print('\nИТОГ ПРОВЕРКИ:', 'ВСЕ ЗНАЧЕНИЯ СХОДЯТСЯ' if fails == 0 else f'{fails} РАСХОЖДЕНИЙ')
