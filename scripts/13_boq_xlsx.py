# -*- coding: utf-8 -*-
"""
Шаг 13. Сводная таблица материалов FTTH по 6 СНП ВКО -> книга Excel.
Листы: «Сводная» (ВОР по материалам), «Параметры сетей», «Методика».
Дизайн — по дизайн-системе навыка xlsx (templates/base.py), палитра professional.
Шрифт Calibri (кириллица, доступен на всех платформах).
"""
import sys, os, json, math

SKILL = '/home/z/my-project/skills/xlsx'
for sub in [SKILL, os.path.join(SKILL, 'templates')]:
    if sub not in sys.path:
        sys.path.insert(0, sub)
import base  # дизайн-токены и фабрики навыка xlsx

# Кириллический контент: Calibri (тонкий шрифт -> допускает полужирный)
base.FONT_NAME = 'Calibri'
base.HEADER_BOLD = True

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = '/home/z/my-project'
OUT = f'{BASE_DIR}/download/Сводная_таблица_материалов_FTTH_ВКО.xlsx'
DATE = '17.09.2026'

D = json.load(open(f'{BASE_DIR}/work/boq_data.json', encoding='utf-8'))
V = D['villages']
VN = [v['name'] for v in V]

# ---------------------------------------------------------------- стили ----
F_INT = '#,##0'
F_KM = '#,##0.0'
F_PCT = '+0.0%;-0.0%;0.0%'

def sect_font():
    return Font(name=base.FONT_NAME, size=11, bold=base.HEADER_BOLD, color=base.PRIMARY)

def sect_fill():
    return PatternFill('solid', fgColor=base.SECONDARY)

def cap_font():
    return Font(name=base.FONT_NAME, size=9, color=base.NEUTRAL_600)

def write_headers(ws, headers, row=4):
    for i, h in enumerate(headers, start=2):
        ws.cell(row=row, column=i, value=h)
    base.style_header_row(ws, row_num=row, col_start=2, col_end=1 + len(headers))

def set_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

# ============================================================ ЛИСТ 1 ========
wb = Workbook()
ws = wb.active
ws.title = 'Сводная'

HDRS = ['№', 'Наименование', 'Ед. изм.', 'Расчёт / норма'] + VN + ['ИТОГО']
LAST = 1 + len(HDRS)          # = 12 (L)
base.setup_sheet(ws, title='Сводная таблица материалов — FTTH (GPON), 6 СНП Восточно-Казахстанской области', last_col=LAST)
ws['B3'] = ('Проект по перечню «СНП ВКО.xlsx», лист «Лоты 12» • сеть спроектирована по спутниковым снимкам '
            'с привязкой к дорогам OSM • допущения — лист «Методика» • подготовлено ' + DATE)
ws['B3'].font = cap_font()
ws.row_dimensions[3].height = 14

write_headers(ws, HDRS)
set_widths(ws, {'A': 3, 'B': 5, 'C': 44, 'D': 9, 'E': 26,
                'F': 13, 'G': 12.5, 'H': 12.5, 'I': 12.5, 'J': 12.5, 'K': 12.5, 'L': 12.5})

SECTIONS = [
    ('А. Оборудование и узлы', [
        ('Шкаф оптический распределительный (ОРШ)', 'шт', '1 на СНП (здание по проекту)', 'orsh', 'int'),
        ('Сплиттер PLC 1×64 (устанавливается в ОРШ)', 'шт', 'ДХ / 64, округление вверх', 'splitters', 'int'),
        ('Пигтейль SC/UPC для кросса ОРШ', 'шт', '1 на ДХ', 'pigtails', 'int'),
        ('Порт PON OLT — справочно (активное оборудование)', 'шт', 'ДХ / 64, округление вверх', 'olt_ports', 'int'),
    ]),
    ('Б. Кабельная продукция (длины с запасом на монтаж 10 %)', [
        ('Кабель оптический самонесущий, 8 волокон', 'км', 'волокна участка = 1,25 × ДХ потока', 'cable_8', 'km'),
        ('Кабель оптический самонесущий, 12 волокон', 'км', 'волокна участка = 1,25 × ДХ потока', 'cable_12', 'km'),
        ('Кабель оптический самонесущий, 16 волокон', 'км', 'волокна участка = 1,25 × ДХ потока', 'cable_16', 'km'),
        ('Кабель оптический самонесущий, 24 волокна', 'км', 'волокна участка = 1,25 × ДХ потока', 'cable_24', 'km'),
        ('Кабель оптический самонесущий, 32 волокна', 'км', 'волокна участка = 1,25 × ДХ потока', 'cable_32', 'km'),
        ('Кабель оптический самонесущий, 48 волокон', 'км', 'волокна участка = 1,25 × ДХ потока', 'cable_48', 'km'),
        ('Кабель оптический самонесущий, 64 волокна', 'км', 'волокна участка = 1,25 × ДХ потока', 'cable_64', 'km'),
        ('Кабель оптический самонесущий, 72 волокна', 'км', 'волокна участка = 1,25 × ДХ потока', 'cable_72', 'km'),
        ('Кабель оптический самонесущий, 96 волокон', 'км', 'волокна участка = 1,25 × ДХ потока', 'cable_96', 'km'),
        ('Кабель дроп-кабельный самонесущий, 2 волокна (1 раб. + 1 рез.)', 'км', 'Σ длин дропов × 1,05', 'drop_cable_km', 'km'),
        ('Справочно: суммарная ёмкость магистрального кабеля', 'волокно-км', 'Σ (длина участка × волокна)', 'fiber_km', 'km'),
    ]),
    ('В. Пассивные узлы', [
        ('Муфта оптическая (проходная / тупиковая)', 'шт', 'по проекту: все ветвления и подключения', 'mufty', 'int'),
    ]),
    ('Г. Материалы для монтажа', [
        ('Бокс абонентский оптический с адаптером SC/UPC', 'шт', '1 на ДХ', 'abonent_boxes', 'int'),
        ('Коннектор оптический механический SC/UPC', 'шт', 'ДХ × 1,1', 'fast_conn', 'int'),
        ('Сварное соединение (оценка объёма работ)', 'шт', '2 × ДХ × 1,1', 'splices', 'int'),
        ('Гильза КДЗС, 60 мм', 'шт', 'по числу сварных соединений', 'kdzs', 'int'),
        ('Комплект подвеса магистрали (кронштейн + спиральный зажим)', 'компл', '30 на 1 км кабеля', 'suspend_kits', 'int'),
        ('Анкерный зажим дроп-кабеля', 'шт', '2 на ДХ', 'drop_anchors', 'int'),
        ('Крепёж дроп-кабеля (скобы / хомуты)', 'шт', '6 на ДХ', 'drop_fix', 'int'),
    ]),
]

row = 5
num = 0
for sect_title, items in SECTIONS:
    # строка секции
    ws.cell(row=row, column=3, value=sect_title)
    for c in range(2, LAST + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = sect_fill()
        cell.font = sect_font()
        cell.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[row].height = 24
    row += 1
    for idx, (name, unit, norma, mkey, kind) in enumerate(items):
        num += 1
        ws.cell(row=row, column=2, value=num)
        ws.cell(row=row, column=3, value=name)
        ws.cell(row=row, column=4, value=unit)
        ws.cell(row=row, column=5, value=norma)
        for vi, v in enumerate(V):
            m = dict(v['materials'])
            m['fiber_km'] = v['fiber_km']
            val = m.get(mkey)
            if val in (0, 0.0):
                val = None
            ws.cell(row=row, column=6 + vi, value=val)
        ws.cell(row=row, column=LAST, value=f'=SUM(F{row}:K{row})')
        # стили
        fill = base.fill_data_row(idx)
        for c in range(2, LAST + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = fill
            cell.font = base.font_body()
            if c == 2:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif c in (3, 5):
                cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            elif c == 4:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = F_KM if kind == 'km' else F_INT
        ws.cell(row=row, column=LAST).font = base.font_subheader()
        ws.row_dimensions[row].height = 24
        row += 1

# примечания
row += 1
notes = [
    'Значения в графах сел — потребность по проекту; графа «ИТОГО» — сумма по шести СНП. Единицы измерения разнородны, итог по столбцам не приводится.',
    'Позиция А.4 (порты PON OLT) — справочно, активное оборудование. Нормы и допущения расчёта — лист «Методика»; параметры сетей — лист «Параметры сетей».',
]
for t in notes:
    ws.cell(row=row, column=2, value=t).font = cap_font()
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=LAST)
    ws.row_dimensions[row].height = 14
    row += 1

ws.freeze_panes = 'F5'
ws.page_setup.orientation = 'landscape'
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.print_title_rows = '4:4'

# ============================================================ ЛИСТ 2 ========
ws2 = wb.create_sheet('Параметры сетей')
H2 = ['№', 'СНП', 'Район / с.о.', 'ДХ (данные заказа)', 'ДХ (проект)', 'Δ к заказу',
      'Муфты', 'Магистраль, км', 'Дропы, км', 'Ср. дроп, м', 'Макс. дроп, м',
      'ОРШ, портов (расч.)', 'Сплиттеры 1:64', 'Верхний участок, волокон', 'Примечание']
L2 = 1 + len(H2)   # = 16 (P)
base.setup_sheet(ws2, title='Параметры спроектированных FTTH-сетей по СНП', last_col=L2)
ws2['B3'] = 'ДХ — домохозяйства. «Верхний участок» — максимальное число волокон магистрали на одном участке (у ОРШ). Подготовлено ' + DATE
ws2['B3'].font = cap_font()
ws2.row_dimensions[3].height = 14
write_headers(ws2, H2)
set_widths(ws2, {'A': 3, 'B': 5, 'C': 17, 'D': 22, 'E': 11, 'F': 10, 'G': 10, 'H': 8,
                 'I': 11, 'J': 9.5, 'K': 9.5, 'L': 10, 'M': 11, 'N': 10, 'O': 11, 'P': 48})

r = 5
for i, v in enumerate(V):
    so = f"{v['raion']} / {v['so']}" if v['so'] != '—' else v['raion']
    note = f"ОРШ: {v['orsh_bld']}. {v['note']}"
    vals = [v['num'], v['name'], so, v['dhx_excel'], v['dhx_served'],
            f'=IFERROR(F{r}/E{r}-1,"")', v['couplers'], v['feeder_km'], v['drop_km'],
            v['avg_drop_m'], v['max_drop_m'], v['orsh_ports'], v['splitters64'], v['top_fibers'], note]
    for ci, val in enumerate(vals, start=2):
        ws2.cell(row=r, column=ci, value=val)
    fill = base.fill_data_row(i)
    for ci in range(2, L2 + 1):
        cell = ws2.cell(row=r, column=ci)
        cell.fill = fill
        cell.font = base.font_body()
        if ci in (2, 4, 7):
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif ci in (3, 16):
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal='right', vertical='center')
    for ci, fmt in [(5, F_INT), (6, F_INT), (7, F_PCT), (8, F_INT), (9, '#,##0.00'),
                    (10, '#,##0.00'), (11, F_INT), (12, F_INT), (13, F_INT), (14, F_INT), (15, F_INT)]:
        ws2.cell(row=r, column=ci).number_format = fmt
    ws2.row_dimensions[r].height = 30
    r += 1

# итоговая строка
tr = r
ws2.cell(row=tr, column=3, value='ИТОГО / справочно')
for ci, f in [(5, '=SUM(E5:E10)'), (6, '=SUM(F5:F10)'), (7, f'=IFERROR(F{tr}/E{tr}-1,"")'),
              (8, '=SUM(H5:H10)'), (9, '=SUM(I5:I10)'), (10, '=SUM(J5:J10)'),
              (11, f'=IFERROR(J{tr}*1000/F{tr},"")'), (14, '=SUM(N5:N10)')]:
    ws2.cell(row=tr, column=ci, value=f)
base.style_total_row(ws2, row_num=tr, col_start=2, col_end=L2)
for ci, fmt in [(5, F_INT), (6, F_INT), (7, F_PCT), (8, F_INT), (9, '#,##0.00'),
                (10, '#,##0.00'), (11, F_INT), (14, F_INT)]:
    c = ws2.cell(row=tr, column=ci)
    c.number_format = fmt
    c.alignment = Alignment(horizontal='right', vertical='center')
ws2.cell(row=tr, column=12, value='ср. по всем').font = Font(name=base.FONT_NAME, size=9, color=base.NEUTRAL_600)
ws2.cell(row=tr, column=12).alignment = Alignment(horizontal='right', vertical='center')

ws2.freeze_panes = 'E5'
ws2.page_setup.orientation = 'landscape'
ws2.page_setup.fitToWidth = 1
ws2.page_setup.fitToHeight = 0

# ============================================================ ЛИСТ 3 ========
ws3 = wb.create_sheet('Методика')
H3 = ['№', 'Положение', 'Принято / описание']
L3 = 1 + len(H3)
base.setup_sheet(ws3, title='Методика расчёта и принятые допущения', last_col=L3)
ws3['B3'] = 'Расчёт выполнен по геометрии спроектированных сетей (ОРШ → муфты → ДХ) без изменения топологии. Подготовлено ' + DATE
ws3['B3'].font = cap_font()
ws3.row_dimensions[3].height = 14
write_headers(ws3, H3)
set_widths(ws3, {'A': 3, 'B': 5, 'C': 34, 'D': 88})

fiber_total = sum(v['fiber_km'] for v in V)
ROWS3 = [
    ('S', '1. Исходные данные', None),
    ('Геометрия сетей', 'Проектные FTTH-сети «ОРШ → муфты → ДХ»: построены по спутниковым снимкам высокого разрешения (~0,38 м/px) с привязкой к дорожной сети OSM; дропы проложены вдоль улиц, дворов и фасадов, все ветвления — в муфтах.'),
    ('Состав сетей', 'Верхнеберезовка — полная сеть села (подтверждена заказчиком, кадр не предоставлялся); Винное — полная сеть на кадре пользователя (550 ДХ, из них 5 за верхней кромкой кадра); Солнечное, Перевальное, Пригородное, Алтайский — в границах предоставленных заказчиком кадров.'),
    ('Число ДХ по данным заказа', '«СНП ВКО.xlsx», лист «Лоты 12»: 3 216 ДХ по шести СНП (Верхнеберезовка 940, Солнечное 366, Перевальное 339, Винное 490, Пригородное 365, Алтайский 716).'),
    ('Обслужено проектом', '2 334 ДХ (см. лист «Параметры сетей»). Отклонения от данных заказа связаны с границами предоставленных кадров, покрытием дорог OSM и фактической застройкой (для Алтайского официальное число ДХ относится ко всему сельскому округу).'),
    ('Длины', 'Длины участков магистрали и дропов — по геометрии проекта с пересчётом пикселей в метры по масштабу мозаик (~0,38 м/px).'),
    ('S', '2. Архитектура', None),
    ('Схема сети', 'GPON, централизованная древовидная топология: все сплиттеры размещаются в ОРШ; каждому ДХ выделяется отдельное волокно от ОРШ до абонентского бокса через цепочку муфт.'),
    ('Сплиттеры', 'PLC 1×64 в ОРШ; количество — ДХ / 64 с округлением вверх. Число портов PON OLT — аналогично (справочно).'),
    ('Дроп-линия', 'Дроп-кабель 2-волоконный (1 рабочее + 1 резервное волокно), самонесущий, от обслуживающей муфты вдоль улицы/двора к зданию.'),
    ('S', '3. Кабельная продукция', None),
    ('Магистральный кабель', 'Самонесущий оптический кабель для воздушной прокладки по существующим опорам; стандартные ёмкости 8 / 12 / 16 / 24 / 32 / 48 / 64 / 72 / 96 волокон.'),
    ('Число волокон на участке', 'ДХ ниже по потоку × 1,25 (резерв 25 %), минимум 8 волокон. При потребности свыше 96 волокон — параллельная прокладка нескольких кабелей меньшей ёмкости (верхние ярусы дерева); допускается замена параллельных кабелей кабелями большей ёмкости по каталогу поставщика.'),
    ('Запас на монтаж', '+10 % к расчётной длине магистрали и +5 % к дроп-кабелю; округление вверх до 0,1 км.'),
    ('Справочно', f'Суммарная расчётная ёмкость магистрали по шести СНП — {fiber_total:,.0f} волокно-км.'.replace(',', ' ')),
    ('S', '4. Узлы и монтажные нормы', None),
    ('Муфты', 'По проекту — все точки ветвления магистрали и подключения групп дропов (1 140 шт по шести СНП). Тип муфты (тупиковая / проходная) и ёмкость по сросткам — по рабочему проекту.'),
    ('ОРШ', '1 на СНП, на общественном/административном здании в центре села (см. лист «Параметры сетей»). Расчётная ёмкость портов — ДХ × 1,1 с округлением вверх до ряда 144 / 288 / 576 / 864 / 1152.'),
    ('Сварные соединения', 'Оценка: 2 на ДХ (сросток «магистраль–дроп» в муфте + пигтейль в кроссе ОРШ) с запасом 10 %. Гильзы КДЗС — по числу сварок.'),
    ('Подвес магистрали', '30 комплектов (кронштейн + спиральный зажим) на 1 км уложенного кабеля; шаг промежуточных креплений ~35 м. Новые опоры не предусмотрены — предполагается использование существующих опорных конструкций.'),
    ('Крепёж дропа', '2 анкерных зажима (у муфты и у стены здания) и 6 промежуточных креплений (скобы/хомуты) на одно ДХ.'),
    ('S', '5. Прочее', None),
    ('Активное оборудование', 'Порты PON OLT приведены справочно; модель, резервирование и комплектацию уточняет заказчик по каталогу поставщика.'),
    ('За рамками расчёта', 'Строительство или аренда опор, заземление, оборудование абонентов (ONT), СМР, накладные и транспортные расходы.'),
]
r = 5
n3 = 0
for item in ROWS3:
    if item[0] == 'S':
        ws3.cell(row=r, column=3, value=item[1])
        for c in range(2, L3 + 1):
            cell = ws3.cell(row=r, column=c)
            cell.fill = sect_fill()
            cell.font = sect_font()
            cell.alignment = Alignment(horizontal='left', vertical='center')
        ws3.row_dimensions[r].height = 24
    else:
        n3 += 1
        name, desc = item
        ws3.cell(row=r, column=2, value=n3)
        ws3.cell(row=r, column=3, value=name)
        ws3.cell(row=r, column=4, value=desc)
        fill = base.fill_data_row(n3)
        for c in range(2, L3 + 1):
            cell = ws3.cell(row=r, column=c)
            cell.fill = fill
            cell.font = base.font_body()
            if c == 2:
                cell.alignment = Alignment(horizontal='center', vertical='top')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        lines = max(math.ceil(len(desc) / 105), math.ceil(len(name) / 30), 1)
        ws3.row_dimensions[r].height = max(22, lines * 14 + 8)
    r += 1

ws3.page_setup.orientation = 'landscape'
ws3.page_setup.fitToWidth = 1
ws3.page_setup.fitToHeight = 0

wb.properties.creator = 'Z.ai'
wb.save(OUT)
print('Сохранено:', OUT)
print('Листы:', wb.sheetnames)
